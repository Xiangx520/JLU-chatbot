"""使用内存中的真实文件验证正文解析，不调用外部服务。"""

import unittest
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject

from JLU_agent.services.RAG.parse_file import FileParseError, FileParser


def make_pdf(pages: list[str | None], password: str | None = None) -> bytes:
    writer = PdfWriter()
    for text in pages:
        page = writer.add_blank_page(width=300, height=300)
        if text is None:
            continue
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})
        })
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 20 250 Td ({text}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    if password is not None:
        writer.encrypt(password)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


class FileParserTests(unittest.TestCase):
    def setUp(self) -> None:
        self.parser = FileParser()

    def test_text_preserves_content_with_and_without_bom(self) -> None:
        text = "  吉林大学\n第二行\n"
        for name in ("history.txt", "history.TXT", "history.md", "history.MD"):
            for encoding in ("utf-8", "utf-8-sig"):
                with self.subTest(name=name, encoding=encoding):
                    self.assertEqual(self.parser.parse_file(text.encode(encoding), name), text)

    def test_markdown_is_not_rendered_or_stripped(self) -> None:
        text = "# 吉林大学\n\n- **办事指南**\n[官网](https://www.jlu.edu.cn)\n"
        self.assertEqual(self.parser.parse_file(text.encode(), "guide.md"), text)

    def test_pdf_keeps_page_order_and_skips_blank_pages(self) -> None:
        self.assertEqual(
            self.parser.parse_file(make_pdf(["First page", None, "Last page"]), "guide.PDF"),
            "First page\nLast page",
        )

    def test_docx_keeps_paragraph_and_table_order(self) -> None:
        document = Document()
        document.add_paragraph("表格之前")
        table = document.add_table(rows=2, cols=2)
        for row, values in zip(table.rows, [("校区", "电话"), ("中心校区", "12345")]):
            for cell, value in zip(row.cells, values):
                cell.text = value
        document.add_paragraph("表格之后")
        document.sections[0].header.paragraphs[0].text = "不读取页眉"
        document.sections[0].footer.paragraphs[0].text = "不读取页脚"
        output = BytesIO()
        document.save(output)
        self.assertEqual(
            self.parser.parse_file(output.getvalue(), "guide.DOCX"),
            "表格之前\n校区\t电话\n中心校区\t12345\n表格之后",
        )

    def test_unsupported_extension(self) -> None:
        for name in ("guide.doc", "guide.xlsx", "guide"):
            with self.subTest(name=name), self.assertRaisesRegex(FileParseError, "不支持"):
                self.parser.parse_file(b"text", name)

    def test_zero_bytes(self) -> None:
        for suffix in ("txt", "md", "pdf", "docx"):
            with self.subTest(suffix=suffix), self.assertRaisesRegex(FileParseError, "文件为空"):
                self.parser.parse_file(b"", f"empty.{suffix}")

    def test_invalid_encoding(self) -> None:
        for suffix in ("txt", "md"):
            with self.subTest(suffix=suffix), self.assertRaisesRegex(FileParseError, "UTF-8"):
                self.parser.parse_file(b"\xff", f"invalid.{suffix}")

    def test_blank_text(self) -> None:
        for suffix in ("txt", "md"):
            with self.subTest(suffix=suffix), self.assertRaisesRegex(FileParseError, "有效正文"):
                self.parser.parse_file(" \n\t".encode("utf-8-sig"), f"blank.{suffix}")

    def test_corrupt_pdf(self) -> None:
        with self.assertRaisesRegex(FileParseError, "无法解析 PDF"):
            self.parser.parse_file(b"%PDF-1.7\ncorrupt", "broken.pdf")

    def test_encrypted_pdf(self) -> None:
        for password in ("secret", ""):
            with self.subTest(password=password), self.assertRaisesRegex(FileParseError, "加密"):
                self.parser.parse_file(make_pdf(["Private"], password), "private.pdf")

    def test_blank_pdf(self) -> None:
        with self.assertRaisesRegex(FileParseError, "有效正文"):
            self.parser.parse_file(make_pdf([None]), "blank.pdf")

    def test_image_only_pdf(self) -> None:
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=300)
        image = DecodedStreamObject()
        image.update({
            NameObject("/Type"): NameObject("/XObject"),
            NameObject("/Subtype"): NameObject("/Image"),
            NameObject("/Width"): NumberObject(1),
            NameObject("/Height"): NumberObject(1),
            NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
            NameObject("/BitsPerComponent"): NumberObject(8),
        })
        image.set_data(b"\xff\xff\xff")
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/XObject"): DictionaryObject({NameObject("/Im0"): writer._add_object(image)})
        })
        stream = DecodedStreamObject()
        stream.set_data(b"q 100 0 0 100 0 0 cm /Im0 Do Q")
        page[NameObject("/Contents")] = writer._add_object(stream)
        output = BytesIO()
        writer.write(output)
        with self.assertRaisesRegex(FileParseError, "OCR"):
            self.parser.parse_file(output.getvalue(), "scan.pdf")

    def test_corrupt_docx_and_missing_package_parts(self) -> None:
        output = BytesIO()
        with ZipFile(output, "w") as archive:
            archive.writestr("unrelated.txt", "not a Word document")
        for data in (b"not a zip", output.getvalue()):
            with self.subTest(data=data), self.assertRaisesRegex(FileParseError, "无法解析 DOCX"):
                self.parser.parse_file(data, "broken.docx")

    def test_docx_with_malformed_xml(self) -> None:
        original = BytesIO()
        Document().save(original)
        output = BytesIO()
        with ZipFile(original) as source, ZipFile(output, "w") as target:
            for name in source.namelist():
                target.writestr(name, b"<broken" if name == "word/document.xml" else source.read(name))
        with self.assertRaisesRegex(FileParseError, "无法解析 DOCX"):
            self.parser.parse_file(output.getvalue(), "broken.docx")

    def test_blank_docx(self) -> None:
        output = BytesIO()
        Document().save(output)
        with self.assertRaisesRegex(FileParseError, "有效正文"):
            self.parser.parse_file(output.getvalue(), "blank.docx")


if __name__ == "__main__":
    unittest.main()
